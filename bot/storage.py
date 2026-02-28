"""Excel storage for user registration and referral tracking."""

from __future__ import annotations

import asyncio
from dataclasses import dataclass
from datetime import datetime, timezone
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.utils import get_column_letter

EXCEL_FILE = Path("data/users.xlsx")
HEADERS = (
    "participant_no",
    "username",
    "registered_at_utc",
    "user_id",
    "referrer_user_id",
    "invites_count",
    "entries_count",
    "referral_qualified",
)
LEGACY_HEADERS = ("full_name", "phone_number", "entered_name")
TECHNICAL_HEADERS = (
    "user_id",
    "referrer_user_id",
    "invites_count",
    "entries_count",
    "referral_qualified",
)

_LOCK = asyncio.Lock()


@dataclass(frozen=True)
class UserRegistration:
    user_id: int
    username: str


@dataclass(frozen=True)
class UserRecord:
    participant_no: int
    user_id: int
    username: str
    registered_at_utc: str
    referrer_user_id: int | None
    invites_count: int
    entries_count: int
    referral_qualified: bool


@dataclass(frozen=True)
class SaveResult:
    record: UserRecord
    is_new_user: bool


def _ensure_file() -> None:
    EXCEL_FILE.parent.mkdir(parents=True, exist_ok=True)
    if EXCEL_FILE.exists():
        return

    wb = Workbook()
    ws = wb.active
    ws.title = "users"
    ws.append(list(HEADERS))
    _hide_technical_columns(ws, _header_map(ws))
    wb.save(EXCEL_FILE)
    wb.close()


def _header_map(ws) -> dict[str, int]:
    return {
        str(cell.value): idx
        for idx, cell in enumerate(ws[1], start=1)
        if cell.value is not None
    }


def _ensure_headers(ws) -> tuple[dict[str, int], bool]:
    existing = _header_map(ws)
    changed = False
    for header in HEADERS:
        if header not in existing:
            ws.cell(row=1, column=ws.max_column + 1, value=header)
            changed = True
    return _header_map(ws), changed


def _drop_legacy_columns(ws, header_map: dict[str, int]) -> bool:
    legacy_indexes = sorted(
        (header_map[h] for h in LEGACY_HEADERS if h in header_map),
        reverse=True,
    )
    changed = False
    for col_idx in legacy_indexes:
        ws.delete_cols(col_idx)
        changed = True
    return changed


def _hide_technical_columns(ws, header_map: dict[str, int]) -> bool:
    changed = False
    for header in TECHNICAL_HEADERS:
        col_idx = header_map.get(header)
        if not col_idx:
            continue
        col_letter = get_column_letter(col_idx)
        dim = ws.column_dimensions[col_letter]
        if not dim.hidden:
            dim.hidden = True
            changed = True
    return changed


def _to_int(value: Any, default: int = 0) -> int:
    try:
        return int(value)
    except (TypeError, ValueError):
        return default


def _to_bool(value: Any, default: bool = False) -> bool:
    if isinstance(value, bool):
        return value
    if value is None:
        return default
    if isinstance(value, (int, float)):
        return bool(value)
    text = str(value).strip().lower()
    return text in {"1", "true", "yes", "y"}


def _next_participant_from_data(ws, header_map: dict[str, int]) -> int:
    p_col = header_map["participant_no"]
    max_no = 0
    for row_idx in range(2, ws.max_row + 1):
        number = _to_int(ws.cell(row=row_idx, column=p_col).value, 0)
        if number > max_no:
            max_no = number
    return max_no + 1


def _migrate_rows(ws, header_map: dict[str, int]) -> bool:
    changed = False
    participant_col = header_map["participant_no"]
    username_col = header_map["username"]
    registered_col = header_map["registered_at_utc"]
    user_id_col = header_map["user_id"]
    referrer_col = header_map["referrer_user_id"]
    invites_col = header_map["invites_count"]
    entries_col = header_map["entries_count"]
    qualified_col = header_map["referral_qualified"]

    next_participant = _next_participant_from_data(ws, header_map)
    now = datetime.now(timezone.utc).replace(microsecond=0).isoformat()

    for row_idx in range(2, ws.max_row + 1):
        user_id = ws.cell(row=row_idx, column=user_id_col).value
        username = ws.cell(row=row_idx, column=username_col).value
        registered_at = ws.cell(row=row_idx, column=registered_col).value
        participant_no = _to_int(ws.cell(row=row_idx, column=participant_col).value, 0)
        invites_count = _to_int(ws.cell(row=row_idx, column=invites_col).value, 0)
        entries_count = _to_int(ws.cell(row=row_idx, column=entries_col).value, 0)
        referral_qualified = _to_bool(ws.cell(row=row_idx, column=qualified_col).value, False)
        referrer = ws.cell(row=row_idx, column=referrer_col).value

        if user_id in (None, "") and username in (None, "") and registered_at in (None, "") and participant_no <= 0:
            continue

        if participant_no <= 0:
            ws.cell(row=row_idx, column=participant_col, value=next_participant)
            next_participant += 1
            changed = True

        if username is None:
            ws.cell(row=row_idx, column=username_col, value="")
            changed = True

        if registered_at in (None, ""):
            ws.cell(row=row_idx, column=registered_col, value=now)
            changed = True

        if invites_count < 0:
            invites_count = 0
            ws.cell(row=row_idx, column=invites_col, value=0)
            changed = True
        elif ws.cell(row=row_idx, column=invites_col).value in (None, ""):
            ws.cell(row=row_idx, column=invites_col, value=0)
            changed = True

        min_entries = 1 + (invites_count // 5)
        if entries_count < min_entries:
            ws.cell(row=row_idx, column=entries_col, value=min_entries)
            changed = True
        elif ws.cell(row=row_idx, column=entries_col).value in (None, ""):
            ws.cell(row=row_idx, column=entries_col, value=min_entries)
            changed = True

        should_be_qualified = invites_count >= 5
        if referral_qualified != should_be_qualified:
            ws.cell(row=row_idx, column=qualified_col, value=should_be_qualified)
            changed = True

        if referrer == "":
            ws.cell(row=row_idx, column=referrer_col, value=None)
            changed = True

    return changed


def _prepare_sheet(ws) -> tuple[dict[str, int], bool]:
    header_map, changed = _ensure_headers(ws)
    if _drop_legacy_columns(ws, header_map):
        changed = True
        header_map = _header_map(ws)
    if _migrate_rows(ws, header_map):
        changed = True
    if _hide_technical_columns(ws, header_map):
        changed = True
    return header_map, changed


def _read_cell(row_values: tuple[Any, ...], idx: int, default: Any = None) -> Any:
    if idx >= len(row_values):
        return default
    value = row_values[idx]
    if value is None:
        return default
    return value


def _find_row(ws, user_id: int, header_map: dict[str, int]) -> int | None:
    user_col = header_map["user_id"] - 1
    for idx, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if _read_cell(row, user_col) == user_id:
            return idx
    return None


def _get_next_participant_no(ws, header_map: dict[str, int]) -> int:
    return _next_participant_from_data(ws, header_map)


def _read_record(ws, row_idx: int, header_map: dict[str, int]) -> UserRecord:
    values = tuple(ws.cell(row=row_idx, column=i).value for i in range(1, ws.max_column + 1))

    def gv(name: str, default: Any = None) -> Any:
        return _read_cell(values, header_map[name] - 1, default)

    ref_raw = gv("referrer_user_id", None)
    return UserRecord(
        participant_no=int(gv("participant_no", 0) or 0),
        user_id=int(gv("user_id", 0) or 0),
        username=str(gv("username", "") or ""),
        registered_at_utc=str(gv("registered_at_utc", "") or ""),
        referrer_user_id=int(ref_raw) if ref_raw not in (None, "") else None,
        invites_count=int(gv("invites_count", 0) or 0),
        entries_count=int(gv("entries_count", 1) or 1),
        referral_qualified=_to_bool(gv("referral_qualified", False), False),
    )


def _save_sync(item: UserRegistration, referrer_user_id: int | None) -> SaveResult:
    _ensure_file()
    wb = load_workbook(EXCEL_FILE)
    ws = wb.active
    header_map, migrated = _prepare_sheet(ws)
    now = datetime.now(timezone.utc).replace(microsecond=0).isoformat()
    dirty = migrated

    row = _find_row(ws, item.user_id, header_map)
    is_new_user = row is None

    if is_new_user:
        participant_no = _get_next_participant_no(ws, header_map)
        if participant_no > 1_000_000:
            wb.close()
            raise ValueError("Participant limit exceeded (1,000,000)")

        row = ws.max_row + 1
        ws.cell(row=row, column=header_map["participant_no"], value=participant_no)
        ws.cell(row=row, column=header_map["username"], value=item.username)
        ws.cell(row=row, column=header_map["registered_at_utc"], value=now)
        ws.cell(row=row, column=header_map["user_id"], value=item.user_id)
        ws.cell(
            row=row,
            column=header_map["referrer_user_id"],
            value=referrer_user_id if referrer_user_id and referrer_user_id != item.user_id else None,
        )
        ws.cell(row=row, column=header_map["invites_count"], value=0)
        ws.cell(row=row, column=header_map["entries_count"], value=1)
        ws.cell(row=row, column=header_map["referral_qualified"], value=False)
        dirty = True
    else:
        existing_username = ws.cell(row=row, column=header_map["username"]).value or ""
        if existing_username != item.username:
            ws.cell(row=row, column=header_map["username"], value=item.username)
            dirty = True

    if is_new_user and referrer_user_id and referrer_user_id != item.user_id:
        ref_row = _find_row(ws, referrer_user_id, header_map)
        if ref_row is not None:
            cur = ws.cell(row=ref_row, column=header_map["invites_count"]).value or 0
            invites = int(cur) + 1
            entries = 1 + (invites // 5)
            ws.cell(row=ref_row, column=header_map["invites_count"], value=invites)
            ws.cell(row=ref_row, column=header_map["entries_count"], value=entries)
            ws.cell(row=ref_row, column=header_map["referral_qualified"], value=(invites >= 5))
            dirty = True

    if dirty:
        wb.save(EXCEL_FILE)
    saved = _read_record(ws, row, header_map)
    wb.close()
    return SaveResult(record=saved, is_new_user=is_new_user)


def _is_registered_sync(user_id: int) -> bool:
    _ensure_file()
    wb = load_workbook(EXCEL_FILE, read_only=False)
    ws = wb.active
    header_map, changed = _prepare_sheet(ws)
    found = _find_row(ws, user_id, header_map) is not None
    if changed:
        wb.save(EXCEL_FILE)
    wb.close()
    return found


def _get_user_record_sync(user_id: int) -> UserRecord | None:
    _ensure_file()
    wb = load_workbook(EXCEL_FILE, read_only=False)
    ws = wb.active
    header_map, changed = _prepare_sheet(ws)
    row = _find_row(ws, user_id, header_map)
    if changed:
        wb.save(EXCEL_FILE)
    if row is None:
        wb.close()
        return None
    record = _read_record(ws, row, header_map)
    wb.close()
    return record


async def save_registration(
    item: UserRegistration, referrer_user_id: int | None = None
) -> SaveResult:
    async with _LOCK:
        return await asyncio.to_thread(_save_sync, item, referrer_user_id)


async def is_registered(user_id: int) -> bool:
    async with _LOCK:
        return await asyncio.to_thread(_is_registered_sync, user_id)


async def get_user_record(user_id: int) -> UserRecord | None:
    async with _LOCK:
        return await asyncio.to_thread(_get_user_record_sync, user_id)


def _get_total_participants_sync() -> int:
    _ensure_file()
    wb = load_workbook(EXCEL_FILE, read_only=False)
    ws = wb.active
    header_map, changed = _prepare_sheet(ws)
    participant_col = header_map["participant_no"] - 1
    total = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        if _read_cell(row, participant_col, None) not in (None, ""):
            total += 1
    if changed:
        wb.save(EXCEL_FILE)
    wb.close()
    return total


async def get_total_participants() -> int:
    async with _LOCK:
        return await asyncio.to_thread(_get_total_participants_sync)
